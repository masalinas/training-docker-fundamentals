"""
theme.py - Módulo de diseño, estilos, colores y componentes visuales para python-pptx.
Cumple con las directrices de diseño corporativo Docker 16:9 y reglas anti-clichés.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Dimensiones 16:9 Widescreen
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Paleta Cromática Docker / Cloud Native
COLOR_PRIMARY_DARK = RGBColor(0x0B, 0x1D, 0x3A)   # Navy Docker profundo
COLOR_PRIMARY_BLUE = RGBColor(0x1D, 0x63, 0xED)   # Azul intenso
COLOR_ACCENT_CYAN  = RGBColor(0x0D, 0xB7, 0xED)   # Cyan Moby Docker
COLOR_TEXT_MAIN    = RGBColor(0x1E, 0x29, 0x3B)   # Slate 800 (texto principal en slides claros)
COLOR_TEXT_MUTED   = RGBColor(0x64, 0x74, 0x8B)   # Slate 500 (subtítulos y detalles)
COLOR_TEXT_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)   # Blanco puro
COLOR_BG_LIGHT     = RGBColor(0xF8, 0xFA, 0xFC)   # Slate 50 (fondo claro de contenido)
COLOR_CARD_BG      = RGBColor(0xFF, 0xFF, 0xFF)   # Fondo blanco para tarjetas
COLOR_CARD_BG_ALT  = RGBColor(0xF1, 0xF5, 0xF9)   # Slate 100 suave
COLOR_CARD_BORDER  = RGBColor(0xCB, 0xD5, 0xE1)   # Slate 300 para bordes
COLOR_CARD_BORDER_ACCENT = RGBColor(0x0D, 0xB7, 0xED)
COLOR_CODE_BG      = RGBColor(0x0F, 0x17, 0x2A)   # Slate 900 (Terminal)
COLOR_CODE_TEXT    = RGBColor(0x38, 0xBD, 0xF8)   # Sky 400 (Terminal prompt/output)
COLOR_CODE_COMMENT = RGBColor(0x94, 0xA3, 0xB8)   # Slate 400
COLOR_BADGE_BG     = RGBColor(0xDC, 0xEA, 0xF8)   # Cyan muy suave
COLOR_BADGE_TEXT   = RGBColor(0x03, 0x4B, 0x82)   # Azul marino
COLOR_ALERT_BG     = RGBColor(0xFE, 0xF2, 0xF2)   # Fondo alerta suave
COLOR_ALERT_BORDER = RGBColor(0xEF, 0x44, 0x44)   # Borde rojo/coral
COLOR_ALERT_TEXT   = RGBColor(0x99, 0x1B, 0x1B)   # Texto rojo oscuro
COLOR_SUCCESS_BG   = RGBColor(0xEC, 0xFD, 0xF5)   # Verde menta
COLOR_SUCCESS_BORDER = RGBColor(0x10, 0xB9, 0x81)

# Tipografías seguras para PowerPoint y LibreOffice
FONT_HEADING = "Calibri"
FONT_BODY    = "Calibri"
FONT_CODE    = "Consolas"


def create_presentation():
    """Inicializa una presentación en blanco configurada a 16:9 panorámico."""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def get_blank_layout(prs):
    """Retorna un layout en blanco (índice 6 en plantillas por defecto)."""
    return prs.slide_layouts[6]


def add_background(slide, color=COLOR_BG_LIGHT):
    """Dibuja un rectángulo de fondo a pantalla completa."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.color.rgb = color
    return bg


def add_title_slide(prs, title, subtitle, module_tag="CURSO TÉCNICO", notes_data=None):
    """Crea una diapositiva de portada o separador de módulo (fondo oscuro)."""
    blank = get_blank_layout(prs)
    slide = prs.slides.add_slide(blank)
    add_background(slide, COLOR_PRIMARY_DARK)

    # Tag superior / Pill
    tag_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.5))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = module_tag.upper()
    p_tag.font.name = FONT_HEADING
    p_tag.font.size = Pt(13)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ACCENT_CYAN

    # Título principal
    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(2.4))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(40)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE

    # Subtítulo / Metadatos
    sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.3), Inches(1.8))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    if notes_data:
        add_speaker_notes(slide, **notes_data)

    return slide


def add_header(slide, title_text, category_tag=""):
    """
    Añade un encabezado limpio y profesional sin líneas decorativas de IA.
    Coloca una categoría en mayúsculas sutil y el título en 28-30pt bold.
    """
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)

    p0 = tf.paragraphs[0]
    if category_tag:
        p0.text = category_tag.upper()
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY_BLUE
        p1 = tf.add_paragraph()
    else:
        p1 = p0

    p1.text = title_text
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_MAIN


def add_card(slide, left, top, width, height, title="", bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    """
    Crea una tarjeta contenedora rectangular con esquinas suavemente redondeadas.
    Retorna la forma para poder operar sobre ella o añadir texto.
    """
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.color.rgb = bg_color

    if title:
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_MAIN

    return card


def add_bullet_list(text_frame, items, font_size=14, text_color=COLOR_TEXT_MAIN, space_after=Pt(8)):
    """
    Añade una lista de viñetas limpia a un text_frame.
    Detecta prefijos en negrita si el ítem contiene ':' (ej. 'Concepto: Explicación').
    """
    for i, item in enumerate(items):
        p = text_frame.add_paragraph() if i > 0 or len(text_frame.text) > 0 else text_frame.paragraphs[0]
        p.space_after = space_after
        p.font.size = Pt(font_size)
        p.font.name = FONT_BODY

        if ":" in item:
            parts = item.split(":", 1)
            run1 = p.add_run()
            run1.text = "• " + parts[0].strip() + ": "
            run1.font.bold = True
            run1.font.color.rgb = text_color

            run2 = p.add_run()
            run2.text = parts[1].strip()
            run2.font.bold = False
            run2.font.color.rgb = text_color
        else:
            run = p.add_run()
            run.text = "• " + item.strip()
            run.font.color.rgb = text_color


def add_code_block(slide, left, top, width, height, code_lines, title="Terminal / Comandos"):
    """
    Dibuja un bloque de terminal con fondo oscuro estilo Slate 900 y fuente Consolas.
    """
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CODE_BG
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1)

    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    if title:
        p_title = tf.paragraphs[0]
        p_title.text = f"$ {title}"
        p_title.font.name = FONT_CODE
        p_title.font.size = Pt(12)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_ACCENT_CYAN
        p_title.space_after = Pt(6)

    for line in code_lines:
        p = tf.add_paragraph()
        p.text = line
        p.font.name = FONT_CODE
        p.font.size = Pt(11)
        if line.startswith("#"):
            p.font.color.rgb = COLOR_CODE_COMMENT
        elif line.startswith("$") or line.startswith(">"):
            p.font.color.rgb = COLOR_CODE_TEXT
            p.font.bold = True
        else:
            p.font.color.rgb = COLOR_TEXT_WHITE
        p.space_after = Pt(2)

    return card


def add_speaker_notes(slide, objective, under_the_hood, classroom_prompt, live_demo):
    """
    Inyecta notas del orador estructuradas en el objeto notes_slide de PowerPoint.
    """
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    
    formatted_notes = (
        f"[OBJETIVO DE LA SLIDE]\n{objective.strip()}\n\n"
        f"[BAJO EL CAPÓ - EXPLICACIÓN TÉCNICA]\n{under_the_hood.strip()}\n\n"
        f"[DINÁMICA DE AULA / PREGUNTA DISPARADORA]\n{classroom_prompt.strip()}\n\n"
        f"[DEMO EN VIVO / COMANDOS]\n{live_demo.strip()}"
    )
    tf.text = formatted_notes
